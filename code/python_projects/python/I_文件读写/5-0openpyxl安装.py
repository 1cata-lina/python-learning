'''pip install openpyxl'''
# Excel 写入操作
from openpyxl import Workbook

# 1. 创建工作簿
wb = Workbook()
# 2. 获取当前工作表
ws = wb.active
# 3. 修改表名
ws.title = "测试数据"

# 方式1：单元格赋值
ws["A1"] = "用户名"
ws["B1"] = "状态码"

# 方式2：行批量写入
ws.append(["admin", 200])
ws.append(["test", 500])
# 写法3：按行列号写入（row行, column列）
ws.cell(row=2, column=1, value="TC001")
ws.cell(row=2, column=2, value="登录模块")
ws.cell(row=2, column=3, value="正确账号密码登录")
ws.cell(row=2, column=4, value="登录成功，跳转首页")
ws.cell(row=2, column=5, value="通过")
# 读取第二行第四列的值
ws.cell(row=2, column=4).value
# 保存文件
import os

# 脚本所在目录 + 文件名 → 绝对路径
save_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "result.xlsx")
wb.save(save_path)
print("已保存到:", save_path)
print("保存成功")