from openpyxl import Workbook, load_workbook

# ========== 题目 1：新建工作簿 + 改表名 + 表头 ==========
wb = Workbook()
ws = wb.active
ws.title = "Monkey异常日志"
ws.append(["时间", "异常类型", "日志内容"])   # 表头
wb.save("monkey_report.xlsx")
print("① 已创建 monkey_report.xlsx")

# ========== 题目 2：append 追加 3 条 ANR 数据 ==========
wb = load_workbook("monkey_report.xlsx")   # 打开已有文件
ws = wb.active
ws.append(["2026-06-18 08:00", "ANR", "主线程阻塞5秒"])
ws.append(["2026-06-18 08:05", "ANR", "输入事件分发超时"])
ws.append(["2026-06-18 08:10", "ANR", "BroadcastReceiver未响应"])
wb.save("monkey_report.xlsx")
print("② 已追加 3 条数据")

# ========== 题目 3：读取，打印所有工作表名 + 每张表标题 ==========
wb = load_workbook("monkey_report.xlsx")
print("③ 所有工作表:", wb.sheetnames)         # ['Monkey异常日志']
for name in wb.sheetnames:
    ws = wb[name]
    title_row = [cell.value for cell in ws[1]]   # 第一行 = 表头
    print(f"  表[{name}] 表头: {title_row}")

# ========== 题目 4：读取第 2 行第 3 列 ==========
wb = load_workbook("monkey_report.xlsx")
ws = wb.active
print("④ 第2行第3列:", ws.cell(row=2, column=3).value)   # 主线程阻塞5秒

# ========== 题目 5：两个工作表 Crash / ANR ==========
wb = Workbook()
ws1 = wb.active
ws1.title = "Crash"
ws1.append(["时间", "异常类型", "日志内容"])

ws2 = wb.create_sheet("ANR")                 # 新建第2张表
ws2.append(["时间", "异常类型", "日志内容"])

wb.save("crash_anr.xlsx")
print("⑤ 已创建:", wb.sheetnames)            # ['Crash', 'ANR']

# ========== 题目 6：只读模式读超大报表 ==========
wb = load_workbook("test_log.xlsx", read_only=True)   # 只读模式，省内存
for sheet in wb.worksheets:
    for row in sheet.iter_rows(values_only=True):     # 每行一个元组
        print(row)
wb.close()                                        # 只读模式需手动关闭
print("⑥ 只读读取完成")

# ========== 题目 7：最大行数 / 最大列数 ==========
wb = load_workbook("monkey_report.xlsx")
ws = wb.active
print("⑦ 最大行数:", ws.max_row)               # 4（表头+3条数据）
print("  最大列数:", ws.max_column)            # 3

# ========== 题目 8：新增工作表放最左侧 + 写数据 ==========
wb = load_workbook("monkey_report.xlsx")
ws = wb.create_sheet("性能数据", 0)            # index=0 → 最左侧
ws.append(["测试项", "耗时(ms)", "测试时间", "测试结果"])
ws.append(["APP冷启动", 1280, "2024-05-20 15:30:25", "通过"])
wb.save("monkey_report.xlsx")
print("⑧ 新表位置:", wb.sheetnames)           # ['性能数据', 'Monkey异常日志']
